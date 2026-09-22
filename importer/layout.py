"""Detects month blocks, weekday/day-number rows, and day columns within a
year sheet (SPEC.md section 7.2). Layout is detected per month block, not
per sheet: the workbook turns out to use at least three header shapes (see
profile.py's output and DECISIONS.md), not a clean early/late split by year.
"""

import calendar
import re
from dataclasses import dataclass
from datetime import date

from openpyxl.worksheet.worksheet import Worksheet

YEAR_SHEET_RE = re.compile(r"^(19|20)\d{2}$")

MONTH_NAMES = [
    "JANUARY",
    "FEBRUARY",
    "MARCH",
    "APRIL",
    "MAY",
    "JUNE",
    "JULY",
    "AUGUST",
    "SEPTEMBER",
    "OCTOBER",
    "NOVEMBER",
    "DECEMBER",
]
MONTH_NUMBER = {name: i + 1 for i, name in enumerate(MONTH_NAMES)}
# The year is optional: later sheets (~2011 onward) label a month block with
# just "January", relying on the sheet's own year tab; a block that carries
# over from the previous year spells the year out explicitly ("DECEMBER
# 2002") specifically because it disagrees with the sheet's own year.
MONTH_LABEL_RE = re.compile(r"^(" + "|".join(MONTH_NAMES) + r")(?:\s+(\d{4}))?$", re.IGNORECASE)

# Weekday-letter tokens as used in the sheet, Monday-first to match
# datetime.date.weekday() (Mon=0 .. Sun=6). "S" is used for both Saturday
# and Sunday; only their position in a run of 7 disambiguates them.
WEEKDAY_LETTERS = ["M", "T", "W", "TR", "F", "S", "S"]
WEEKDAY_TOKEN_SET = set(WEEKDAY_LETTERS)

MAX_DAY_COLUMNS = 32  # a month has at most 31 days; +1 guards off-by-one scans


def _is_weekday_token(value: object) -> bool:
    return isinstance(value, str) and value.strip().upper() in WEEKDAY_TOKEN_SET


def _is_dayish(value: object) -> bool:
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return True
    return isinstance(value, str) and value.strip().startswith("=")


def _find_day_number_run(ws: Worksheet, row: int, max_col: int = 45) -> tuple[int, int] | None:
    """Finds a run of >= 3 columns in `row` holding ascending day numbers
    1, 2, 3, ... starting at some column. Early sheets express days 2+ as
    `=SUM(prev+1)` formulas rather than literal ints (data_only=False), so a
    formula cell is accepted on faith once the run has already started on a
    literal `1` — evaluating it would require a second, data_only load."""
    col = 2
    while col <= max_col:
        if not _is_dayish(ws.cell(row=row, column=col).value):
            col += 1
            continue
        start = col
        n = 0
        c = col
        expect = 1
        while c <= max_col:
            v = ws.cell(row=row, column=c).value
            ok = (isinstance(v, (int, float)) and not isinstance(v, bool) and int(v) == expect) or (
                n > 0 and isinstance(v, str) and v.strip().startswith("=")
            )
            if not ok:
                break
            n += 1
            expect += 1
            c += 1
        if n >= 3:
            return start, n
        col = c + 1 if c > col else col + 1
    return None


def _weekday_hits(ws: Worksheet, row: int, first_col: int, length: int) -> int:
    return sum(
        1
        for c in range(first_col, first_col + length)
        if _is_weekday_token(ws.cell(row=row, column=c).value)
    )


@dataclass(frozen=True)
class MonthHeader:
    label: str
    year: int
    month: int
    label_row: int
    weekday_row: int
    daynum_row: int
    first_day_col: int
    first_berth_row: int

    def days_in_month(self) -> int:
        return calendar.monthrange(self.year, self.month)[1]

    def column_for_day(self, day: int) -> int:
        return self.first_day_col + (day - 1)

    def day_for_column(self, col: int) -> int | None:
        day = col - self.first_day_col + 1
        if 1 <= day <= self.days_in_month():
            return day
        return None

    def date_for_column(self, col: int) -> date | None:
        day = self.day_for_column(col)
        if day is None:
            return None
        return date(self.year, self.month, day)


def find_month_headers(ws: Worksheet, default_year: int | None = None) -> list[MonthHeader]:
    """Finds every month block in `ws`, trying each of the header shapes
    observed in the workbook, in order, per label row:

    A. day numbers on the label row itself, weekday letters the row after
       (the shape SPEC.md section 7.2 calls "later").
    B. weekday letters on the label row itself, day numbers the row after.
    C. label row alone, weekday letters the next row, day numbers the row
       after that (SPEC.md section 7.2's "early" shape).

    A header whose weekday/day-number rows can't be confidently located
    still gets recorded, with a best-effort guess, so callers can flag it
    rather than silently skipping a month block.

    `default_year` is used when the label omits a year (e.g. "January" on
    later sheets, which rely on the sheet's own year tab) — pass the year
    sheet's name, parsed as an int. A block whose label spells out a year
    ("DECEMBER 2002") always uses that instead, since it exists precisely
    because it disagrees with the sheet's own year (a cross-year carry-over
    block at the top of the sheet).
    """
    headers: list[MonthHeader] = []
    for row in range(1, ws.max_row + 1):
        label = ws.cell(row=row, column=1).value
        if not isinstance(label, str):
            continue
        m = MONTH_LABEL_RE.match(label.strip())
        if not m:
            continue
        month = MONTH_NUMBER[m.group(1).upper()]
        if m.group(2) is not None:
            year = int(m.group(2))
        elif default_year is not None:
            year = default_year
        else:
            continue

        run = _find_day_number_run(ws, row)
        if run is not None:
            first_col, length = run
            if _weekday_hits(ws, row + 1, first_col, length) >= length - 1:
                headers.append(
                    MonthHeader(label.strip(), year, month, row, row + 1, row, first_col, row + 2)
                )
                continue

        run = _find_day_number_run(ws, row + 1)
        if run is not None:
            first_col, length = run
            if _weekday_hits(ws, row, first_col, length) >= length - 1:
                headers.append(
                    MonthHeader(label.strip(), year, month, row, row, row + 1, first_col, row + 2)
                )
                continue

        run = _find_day_number_run(ws, row + 2)
        if run is not None:
            first_col, length = run
            if _weekday_hits(ws, row + 1, first_col, length) >= length - 1:
                headers.append(
                    MonthHeader(
                        label.strip(), year, month, row, row + 1, row + 2, first_col, row + 3
                    )
                )
                continue

        headers.append(MonthHeader(label.strip(), year, month, row, row + 1, row + 1, 2, row + 2))

    return headers


def is_weekday_row(ws: Worksheet, row: int, max_col: int = MAX_DAY_COLUMNS + 13) -> bool:
    """True if `row` looks like a weekday-letter row. Used to detect where
    a berth block ends (the run of berth rows stops at the next weekday
    row, a blank row, or the sheet edge)."""
    hits = sum(1 for c in range(2, max_col) if _is_weekday_token(ws.cell(row=row, column=c).value))
    return hits >= 5


def calendar_mismatches(ws: Worksheet, header: MonthHeader) -> list[tuple[int, str, str]]:
    """Compares the weekday-letter row against the real calendar for
    `header`'s year/month (SPEC.md section 7.2's sanity check). Returns a
    list of (column, expected_letter, actual_value) for every day column
    that doesn't match, empty if the row matches perfectly.
    """
    mismatches: list[tuple[int, str, str]] = []
    for day in range(1, header.days_in_month() + 1):
        col = header.column_for_day(day)
        expected = WEEKDAY_LETTERS[date(header.year, header.month, day).weekday()]
        actual = ws.cell(row=header.weekday_row, column=col).value
        actual_str = actual.strip().upper() if isinstance(actual, str) else actual
        if actual_str != expected:
            mismatches.append((col, expected, str(actual)))
    return mismatches
