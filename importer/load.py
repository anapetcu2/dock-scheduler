"""Loads the workbook into the database (SPEC.md section 7.6-7.7). Writes
directly with SQLAlchemy rather than through app/services/bookings.py,
because historical data is expected to violate the booking rules (that's
the whole point of the importer) and must still be imported.
"""

import re
from dataclasses import dataclass, field
from datetime import date

import openpyxl
from openpyxl.worksheet.worksheet import Worksheet
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models.berths import Berth
from app.models.bookings import Booking, BookingKind, BookingSource, BookingStatus
from app.models.contacts import Contact, Organization
from app.models.import_issues import ImportIssue, ImportIssueSeverity, ImportIssueType
from app.models.vessels import Vessel, VesselContact
from importer.classify import classify
from importer.layout import YEAR_SHEET_RE, MonthHeader, calendar_mismatches, find_month_headers
from importer.normalize import VesselGroup, VesselMention, group_vessel_mentions
from importer.ruleset import ImporterConfig, load_config
from importer.spans import (
    Span,
    SpanIssue,
    berth_rows_for_block,
    build_merge_lookup,
    extract_spans,
    find_unlabeled_entries,
)
from importer.vessels import CONTACT_DISCLAIMER, VesselDetailRecord, parse_vessel_detail_tab

BERTH_LABEL_RE = re.compile(r"^(.*?)\s*-\s*(\d+(?:\.\d+)?)'$")
IGNORED_SHEETS = {"Tours"}
SUMMARY_SHEET = "8YR Dock Summary"
SCIENCE_SHEET = "Science"
YACHTS_SHEET = "Yachts"


@dataclass
class RawIssue:
    issue_type: ImportIssueType
    severity: ImportIssueSeverity
    message: str
    source_ref: str | None = None
    details: dict | None = None
    entity_type: str | None = None
    entity_id: int | None = None


@dataclass
class RawBooking:
    berth_label: str
    kind: str  # "vessel" | "event" | "closure"
    start_date: date
    end_date: date
    vessel_mention: VesselMention | None
    title: str | None
    notes: str | None
    source_ref: str


@dataclass
class ImportStats:
    issues: list[RawIssue] = field(default_factory=list)
    bookings: list[RawBooking] = field(default_factory=list)
    berth_names: dict[str, tuple[str, float | None]] = field(
        default_factory=dict
    )  # label -> (name, length)
    vessel_mentions: list[VesselMention] = field(default_factory=list)
    vessel_detail_records: list[VesselDetailRecord] = field(default_factory=list)
    classification_counts: dict[str, int] = field(default_factory=dict)
    year_berth_days: dict[tuple[int, str], int] = field(
        default_factory=dict
    )  # (year, berth name) -> days
    summary_sheet_days: dict[tuple[int, str], int] = field(default_factory=dict)


def parse_berth_label(label: str, config: ImporterConfig) -> tuple[str, float | None]:
    if label in config.unlengthed_berths:
        return label.rstrip(":").strip(), None
    m = BERTH_LABEL_RE.match(label)
    if m:
        return m.group(1).strip(), float(m.group(2))
    return label.strip(), None


def _add_span_days(stats: ImportStats, year: int, berth_name: str, days: int) -> None:
    key = (year, berth_name)
    stats.year_berth_days[key] = stats.year_berth_days.get(key, 0) + days


def _process_month_block(
    ws: Worksheet,
    header: MonthHeader,
    sheet_name: str,
    config: ImporterConfig,
    merge_lookup: dict,
    stats: ImportStats,
) -> dict[str, tuple[list[Span], list[SpanIssue]]]:
    """Returns {berth_label: (spans, issues)} for this block, and records
    berth labels / unlabeled entries / calendar mismatches as a side
    effect. Cross-month joining happens later, across the returned dict of
    consecutive blocks."""
    mismatches = calendar_mismatches(ws, header)
    if mismatches:
        stats.issues.append(
            RawIssue(
                ImportIssueType.LAYOUT_MISMATCH,
                ImportIssueSeverity.warning,
                f'"{header.label}" on {sheet_name}: weekday row does not match the real '
                f"calendar for {header.year}-{header.month:02d} ({len(mismatches)} day(s) off).",
                source_ref=f"{sheet_name}!row{header.weekday_row}",
                details={
                    "mismatches": [{"col": c, "expected": e, "actual": a} for c, e, a in mismatches]
                },
            )
        )

    berth_rows = berth_rows_for_block(ws, header)
    per_berth: dict[str, tuple[list[Span], list[SpanIssue]]] = {}

    for row, label in berth_rows:
        if label not in stats.berth_names:
            stats.berth_names[label] = parse_berth_label(label, config)

        spans, issues = extract_spans(ws, header, row, label, sheet_name, config, merge_lookup)
        per_berth[label] = (spans, issues)

    for entry in find_unlabeled_entries(ws, header, berth_rows, sheet_name):
        stats.issues.append(
            RawIssue(
                ImportIssueType.UNLABELED_ROW_ENTRY,
                ImportIssueSeverity.warning,
                f'Text {entry.text!r} found below "{entry.berth_above}" with no berth label '
                "of its own.",
                source_ref=entry.source_ref,
            )
        )

    return per_berth


def _join_across_blocks(
    prev_block: dict[str, tuple[list[Span], list[SpanIssue]]],
    prev_header: MonthHeader,
    curr_block: dict[str, tuple[list[Span], list[SpanIssue]]],
    curr_header: MonthHeader,
) -> None:
    """Mutates `curr_block` in place: SPEC.md section 7.3 rule 6. A span
    that reaches the last day of `prev_header`'s month is merged with
    `curr_header`'s day-1 span for the same berth if they share a
    normalized name, or if day 1 was an unnamed run (which would otherwise
    be an ORPHAN_FILL) that has nothing else claiming it."""
    for label, (curr_spans, curr_issues) in curr_block.items():
        if label not in prev_block:
            continue
        prev_spans, _ = prev_block[label]
        prev_last_day_spans = [
            s
            for s in prev_spans
            if s.end_col == prev_header.column_for_day(prev_header.days_in_month())
        ]
        if not prev_last_day_spans:
            continue
        prev_span = prev_last_day_spans[-1]

        day1_col = curr_header.column_for_day(1)
        matching_span = next((s for s in curr_spans if s.start_col == day1_col), None)
        if matching_span is not None:
            if matching_span.text.strip().upper() == prev_span.text.strip().upper():
                prev_span.end_date_override = curr_header.date_for_column(matching_span.end_col)
                prev_span.source_ref += f" + {matching_span.source_ref}"
                curr_spans.remove(matching_span)
            continue

        orphan = next(
            (
                i
                for i in curr_issues
                if i.code == "ORPHAN_FILL" and i.berth_label == label and i.start_col == day1_col
            ),
            None,
        )
        if orphan is not None:
            prev_span.end_date_override = curr_header.date_for_column(orphan.end_col)
            prev_span.source_ref += f" + {orphan.source_ref}"
            curr_issues.remove(orphan)


def _span_dates(header: MonthHeader, span: Span) -> tuple[date, date]:
    start = header.date_for_column(span.start_col)
    end = (
        span.end_date_override
        if span.end_date_override is not None
        else header.date_for_column(span.end_col)
    )
    assert start is not None and end is not None
    return start, end


def _classify_block(
    header: MonthHeader,
    per_berth: dict[str, tuple[list[Span], list[SpanIssue]]],
    sheet_name: str,
    config: ImporterConfig,
    stats: ImportStats,
) -> None:
    for label, (spans, issues) in per_berth.items():
        for issue in issues:
            if issue.code == "ORPHAN_FILL":
                stats.issues.append(
                    RawIssue(
                        ImportIssueType.ORPHAN_FILL,
                        ImportIssueSeverity.warning,
                        issue.message,
                        issue.source_ref,
                    )
                )
            elif issue.code == "AMBIGUOUS_BOUNDARY":
                stats.issues.append(
                    RawIssue(
                        ImportIssueType.AMBIGUOUS_BOUNDARY,
                        ImportIssueSeverity.warning,
                        issue.message,
                        issue.source_ref,
                    )
                )

        occupying: list[tuple[Span, date, date, str]] = []  # span, start, end, kind
        notes_spans: list[tuple[Span, date, date]] = []

        for span in spans:
            start, end = _span_dates(header, span)
            berth_name, _ = stats.berth_names[label]

            result = classify(span.text, config)
            stats.classification_counts[result.kind] = (
                stats.classification_counts.get(result.kind, 0) + 1
            )

            if result.kind == "vessel":
                # The 8YR Dock Summary tab's "days booked" turns out to
                # count vessel occupancy only — a 31-day December closure
                # spot-checked against a "days" value of 1 in the summary
                # confirmed this (see DECISIONS.md, Phase 4). Events and
                # closures are excluded from this comparison accordingly.
                _add_span_days(stats, start.year, berth_name, (end - start).days + 1)

            if result.kind == "note":
                notes_spans.append((span, start, end))
                continue
            if result.kind == "vessel":
                stats.vessel_mentions.append(result.vessel)
                stats.bookings.append(
                    RawBooking(
                        label, "vessel", start, end, result.vessel, None, None, span.source_ref
                    )
                )
            else:  # closure | event
                stats.bookings.append(
                    RawBooking(
                        label, result.kind, start, end, None, span.text, None, span.source_ref
                    )
                )
            occupying.append((span, start, end, result.kind))

        for note_span, n_start, n_end in notes_spans:
            attached = False
            for booking in stats.bookings:
                if booking.berth_label != label:
                    continue
                if booking.start_date <= n_end and n_start <= booking.end_date:
                    booking.notes = (
                        f"{booking.notes}; {note_span.text}" if booking.notes else note_span.text
                    )
                    attached = True
                    break
            if not attached:
                stats.issues.append(
                    RawIssue(
                        ImportIssueType.UNATTACHED_NOTE,
                        ImportIssueSeverity.info,
                        f'Note {note_span.text!r} on "{label}" does not overlap any booking.',
                        note_span.source_ref,
                    )
                )


def scan_workbook(path, config: ImporterConfig) -> ImportStats:
    wb = openpyxl.load_workbook(path, data_only=False)
    stats = ImportStats()

    for sheet_name in wb.sheetnames:
        if sheet_name in IGNORED_SHEETS or sheet_name in (
            SUMMARY_SHEET,
            SCIENCE_SHEET,
            YACHTS_SHEET,
        ):
            continue
        if not YEAR_SHEET_RE.match(sheet_name):
            continue

        ws = wb[sheet_name]
        merge_lookup = build_merge_lookup(ws)
        headers = find_month_headers(ws, default_year=int(sheet_name))

        blocks = [
            _process_month_block(ws, h, sheet_name, config, merge_lookup, stats) for h in headers
        ]

        for i in range(1, len(blocks)):
            prev_header, curr_header = headers[i - 1], headers[i]
            is_consecutive = (curr_header.year, curr_header.month) == (
                (prev_header.year, prev_header.month + 1)
                if prev_header.month < 12
                else (prev_header.year + 1, 1)
            )
            if is_consecutive:
                _join_across_blocks(blocks[i - 1], prev_header, blocks[i], curr_header)

        for header, per_berth in zip(headers, blocks, strict=True):
            _classify_block(header, per_berth, sheet_name, config, stats)

    for label in config.unlengthed_berths:
        if label in stats.berth_names:
            name, _ = stats.berth_names[label]
            stats.issues.append(
                RawIssue(
                    ImportIssueType.UNKNOWN_BERTH_LENGTH,
                    ImportIssueSeverity.info,
                    f'"{name}" has no length in the workbook; created with length_ft = null.',
                )
            )

    for sheet_name in (SCIENCE_SHEET, YACHTS_SHEET):
        if sheet_name in wb.sheetnames:
            records = parse_vessel_detail_tab(wb[sheet_name], sheet_name, config)
            stats.vessel_detail_records.extend(records)
            for r in records:
                if r.conflicting_length is not None:
                    used, discarded = r.conflicting_length
                    stats.issues.append(
                        RawIssue(
                            ImportIssueType.CONFLICTING_VESSEL_LENGTH,
                            ImportIssueSeverity.warning,
                            f"{r.mention.prefix} {r.mention.name}: "
                            f"LOA: says {used}', name cell says {discarded}'; using {used}'.",
                            source_ref=r.source_ref,
                        )
                    )

    if SUMMARY_SHEET in wb.sheetnames:
        _check_summary_only_berths(wb[SUMMARY_SHEET], config, stats)
        stats.summary_sheet_days = parse_summary_sheet(wb[SUMMARY_SHEET])

    return stats


def parse_summary_sheet(ws: Worksheet) -> dict[tuple[int, str], int]:
    """Reads the 8YR Dock Summary tab into {(year, berth_name): days}, for
    SPEC.md section 7.8's comparison against the importer's own count."""
    year_cols = {}
    for cell in ws[1]:
        if isinstance(cell.value, int):
            year_cols[cell.column] = cell.value

    result: dict[tuple[int, str], int] = {}
    for row in ws.iter_rows(min_row=2):
        name = row[0].value
        if not isinstance(name, str) or name.strip().lower() == "total days":
            continue
        for cell in row:
            if cell.column in year_cols and isinstance(cell.value, (int, float)):
                result[(year_cols[cell.column], name.strip())] = int(cell.value)
    return result


def _check_summary_only_berths(ws: Worksheet, config: ImporterConfig, stats: ImportStats) -> None:
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        name = row[0].value
        if isinstance(name, str) and name.strip() in config.summary_only_berths:
            stats.issues.append(
                RawIssue(
                    ImportIssueType.UNKNOWN_BERTH,
                    ImportIssueSeverity.info,
                    f'"{name.strip()}" appears on the 8YR Dock Summary tab but not in any year '
                    "grid; no berth created for it.",
                    source_ref=f"{SUMMARY_SHEET}!{row[0].coordinate}",
                )
            )


def write_to_db(session: Session, stats: ImportStats, config: ImporterConfig) -> dict:
    """Writes berths, vessels (+ orgs/contacts), bookings, and import
    issues. Returns summary counts for the report.

    Berths, vessels, organizations and contacts have no `source` column
    (only bookings and import_issues do — see SPEC.md section 5), so
    `--reset` never deletes them (it could delete rows created through the
    live app, not just imported ones). For `--reset` to be idempotent
    without violating their unique constraints, this looks each one up by
    its natural key (name / normalized_key) before creating it, reusing an
    existing row from a prior run instead of colliding with it.
    """
    berths_by_name: dict[str, Berth] = {}
    for _label, (name, length) in stats.berth_names.items():
        if name in berths_by_name:
            continue
        existing = session.scalar(select(Berth).where(Berth.name == name))
        if existing is not None:
            berths_by_name[name] = existing
            continue
        berth = Berth(name=name, length_ft=length, is_active=True, sort_order=len(berths_by_name))
        session.add(berth)
        berths_by_name[name] = berth
    session.flush()

    all_mentions = list(stats.vessel_mentions) + [r.mention for r in stats.vessel_detail_records]
    groups = group_vessel_mentions(all_mentions, config)

    detail_by_key: dict[str, VesselDetailRecord] = {}
    for record in stats.vessel_detail_records:
        from app.services.vessels import normalize_vessel_key

        key = normalize_vessel_key(record.mention.prefix, record.mention.name)
        detail_by_key[key] = record

    organizations_by_name: dict[str, Organization] = {}
    vessels_by_key: dict[str, Vessel] = {}

    for key, group in groups.items():
        _create_vessel(
            session,
            key,
            group,
            detail_by_key.get(key),
            organizations_by_name,
            vessels_by_key,
            stats,
        )

    for group in groups.values():
        if len(group.variants) > 1:
            stats.issues.append(
                RawIssue(
                    ImportIssueType.NAME_VARIANTS_MERGED,
                    ImportIssueSeverity.info,
                    f"Merged {len(group.variants)} name variants into "
                    f"{group.prefix} {group.display_name}: " + ", ".join(group.variants),
                )
            )

    session.flush()
    _write_bookings(session, stats, berths_by_name, vessels_by_key, config)

    for issue in stats.issues:
        session.add(
            ImportIssue(
                issue_type=issue.issue_type,
                severity=issue.severity,
                message=issue.message,
                source_ref=issue.source_ref,
                details=issue.details,
                entity_type=issue.entity_type,
                entity_id=issue.entity_id,
            )
        )

    return {
        "berths": len(berths_by_name),
        "vessels": len(vessels_by_key),
        "bookings": len(stats.bookings),
        "issues": len(stats.issues),
    }


def _create_vessel(
    session: Session,
    key: str,
    group: VesselGroup,
    detail: VesselDetailRecord | None,
    organizations_by_name: dict[str, Organization],
    vessels_by_key: dict[str, Vessel],
    stats: ImportStats,
) -> None:
    notes = None
    if detail is not None and (detail.captain_names or detail.phones or detail.emails):
        parts = [CONTACT_DISCLAIMER]
        if detail.captain_names:
            parts.append("Contacts: " + ", ".join(dict.fromkeys(detail.captain_names)))
        if detail.phones:
            parts.append("Phones: " + ", ".join(dict.fromkeys(detail.phones)))
        if detail.emails:
            parts.append("Emails: " + ", ".join(dict.fromkeys(detail.emails)))
        notes = " ".join(parts)

    organization_id = None
    if detail is not None and detail.organization_name:
        org = organizations_by_name.get(detail.organization_name)
        if org is None:
            org = session.scalar(
                select(Organization).where(Organization.name == detail.organization_name)
            )
        if org is None:
            org = Organization(name=detail.organization_name)
            session.add(org)
            session.flush()
        organizations_by_name[detail.organization_name] = org
        organization_id = org.id

    existing = session.scalar(select(Vessel).where(Vessel.normalized_key == key))
    if existing is not None:
        vessels_by_key[key] = existing
        return

    vessel = Vessel(
        name=group.display_name,
        type_prefix=group.prefix,
        normalized_key=key,
        loa_ft=detail.loa_ft if detail else None,
        draft_ft=detail.draft_ft if detail else None,
        organization_id=organization_id,
        is_active=True,
        notes=notes,
    )
    session.add(vessel)
    session.flush()
    vessels_by_key[key] = vessel

    if detail is not None:
        for captain in dict.fromkeys(detail.captain_names):
            contact = Contact(name=captain, organization_id=organization_id)
            session.add(contact)
            session.flush()
            session.add(VesselContact(vessel_id=vessel.id, contact_id=contact.id, role="captain"))


def _write_bookings(
    session: Session,
    stats: ImportStats,
    berths_by_name: dict[str, Berth],
    vessels_by_key: dict[str, Vessel],
    config: ImporterConfig,
) -> None:
    from app.services.vessels import normalize_vessel_key

    by_berth: dict[str, list[RawBooking]] = {}
    for booking in stats.bookings:
        name, _ = stats.berth_names[booking.berth_label]
        by_berth.setdefault(name, []).append(booking)

    for berth_name, bookings in by_berth.items():
        berth = berths_by_name[berth_name]
        bookings.sort(key=lambda b: (b.start_date, b.end_date))
        active: list[Booking] = []

        for raw in bookings:
            vessel_id = None
            if raw.vessel_mention is not None:
                key = normalize_vessel_key(raw.vessel_mention.prefix, raw.vessel_mention.name)
                vessel = vessels_by_key.get(key)
                vessel_id = vessel.id if vessel else None

            overlaps = [
                b for b in active if b.start_date <= raw.end_date and raw.start_date <= b.end_date
            ]
            status = BookingStatus.legacy_conflict if overlaps else BookingStatus.confirmed

            booking = Booking(
                berth_id=berth.id,
                kind=BookingKind(raw.kind),
                vessel_id=vessel_id,
                title=raw.title,
                start_date=raw.start_date,
                end_date=raw.end_date,
                status=status,
                notes=raw.notes,
                source=BookingSource.import_,
                source_ref=raw.source_ref,
                created_by=None,
            )
            session.add(booking)
            session.flush()
            active.append(booking)

            if overlaps:
                for other in overlaps:
                    stats.issues.append(
                        RawIssue(
                            ImportIssueType.HISTORICAL_OVERLAP,
                            ImportIssueSeverity.warning,
                            f'"{berth_name}" {raw.start_date}–{raw.end_date} overlaps '
                            f"booking #{other.id} ({other.start_date}–{other.end_date}).",
                            source_ref=raw.source_ref,
                            entity_type="booking",
                            entity_id=booking.id,
                        )
                    )


def run_import(
    session: Session, path, *, reset: bool, config: ImporterConfig | None = None
) -> tuple[ImportStats, dict]:
    config = config or load_config()
    if reset:
        session.query(Booking).filter(Booking.source == BookingSource.import_).delete()
        session.query(ImportIssue).delete()
        session.flush()

    stats = scan_workbook(path, config)
    summary = write_to_db(session, stats, config)
    return stats, summary
