"""Reconstructs booking spans from the cells in a berth row, within one
month block (SPEC.md section 7.3).
"""

from dataclasses import dataclass
from datetime import date

from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from importer.fills import is_background_fill
from importer.layout import MonthHeader, is_weekday_row
from importer.ruleset import ImporterConfig


def cell_ref(row: int, col: int) -> str:
    return f"{get_column_letter(col)}{row}"


@dataclass
class Span:
    berth_label: str
    start_col: int
    end_col: int
    text: str
    source_ref: str
    # Set only when a cross-month join (SPEC.md section 7.3 rule 6) extends
    # this span into a later month block, whose day columns don't share an
    # origin with this span's own header (a later month's calendar-aligned
    # start column can differ arbitrarily from this one's). `end_col` alone
    # can't express that; this holds the real date instead.
    end_date_override: date | None = None


@dataclass
class SpanIssue:
    code: str  # "ORPHAN_FILL" | "AMBIGUOUS_BOUNDARY"
    message: str
    source_ref: str
    berth_label: str = ""
    # Set for ORPHAN_FILL only, so a caller doing cross-month joins (SPEC.md
    # section 7.3 rule 6) can recover the exact extent of an unnamed run
    # that turned out to continue the previous month's span, without
    # re-parsing `source_ref`.
    start_col: int | None = None
    end_col: int | None = None


def build_merge_lookup(ws: Worksheet) -> dict[tuple[int, int], object]:
    """Maps every (row, col) covered by a merged range to that range, so a
    cell's membership can be checked in O(1) instead of scanning all ranges
    per cell."""
    lookup: dict[tuple[int, int], object] = {}
    for rng in ws.merged_cells.ranges:
        for row in range(rng.min_row, rng.max_row + 1):
            for col in range(rng.min_col, rng.max_col + 1):
                lookup[(row, col)] = rng
    return lookup


def _is_booking_cell(ws: Worksheet, row: int, col: int, config: ImporterConfig) -> bool:
    cell = ws.cell(row=row, column=col)
    if isinstance(cell.value, str) and cell.value.strip():
        return True
    return not is_background_fill(cell, config.background_fills)


def _spans_from_run(
    run_cells: list[tuple[int, str | None]],  # (col, text_or_None)
    berth_label: str,
    sheet_name: str,
    row: int,
) -> tuple[list[Span], list[SpanIssue]]:
    spans: list[Span] = []
    issues: list[SpanIssue] = []

    named = [i for i, (_, text) in enumerate(run_cells) if text is not None]
    if not named:
        first_col, last_col = run_cells[0][0], run_cells[-1][0]
        issues.append(
            SpanIssue(
                "ORPHAN_FILL",
                f'Colored/filled run with no name on "{berth_label}"',
                f"{sheet_name}!{cell_ref(row, first_col)}:{cell_ref(row, last_col)}",
                berth_label=berth_label,
                start_col=first_col,
                end_col=last_col,
            )
        )
        return spans, issues

    for i, name_idx in enumerate(named):
        text = run_cells[name_idx][1]
        assert text is not None
        # Rule 3: unnamed cells before the first name belong to that span;
        # unnamed cells after a name belong to it until the next name (or
        # run end). Both fall out of this start/end formula.
        start_col = run_cells[0][0] if i == 0 else run_cells[name_idx][0]
        if i + 1 < len(named):
            next_idx = named[i + 1]
            end_col = run_cells[next_idx][0] - 1
            if next_idx == name_idx + 1:
                # Rule 5: two names with no gap between them.
                next_text = run_cells[next_idx][1]
                issues.append(
                    SpanIssue(
                        "AMBIGUOUS_BOUNDARY",
                        f'Two names back-to-back on "{berth_label}": {text!r} then {next_text!r}',
                        f"{sheet_name}!{cell_ref(row, run_cells[name_idx][0])}",
                    )
                )
        else:
            end_col = run_cells[-1][0]
        spans.append(
            Span(
                berth_label=berth_label,
                start_col=start_col,
                end_col=end_col,
                text=text,
                source_ref=f"{sheet_name}!{cell_ref(row, start_col)}:{cell_ref(row, end_col)}",
            )
        )
    return spans, issues


def extract_spans(
    ws: Worksheet,
    header: MonthHeader,
    berth_row: int,
    berth_label: str,
    sheet_name: str,
    config: ImporterConfig,
    merge_lookup: dict[tuple[int, int], object],
) -> tuple[list[Span], list[SpanIssue]]:
    spans: list[Span] = []
    issues: list[SpanIssue] = []
    days = header.days_in_month()

    day = 1
    while day <= days:
        col = header.column_for_day(day)
        merge_rng = merge_lookup.get((berth_row, col))

        if merge_rng is not None:
            # Rule 2: a merged range is always its own span, processed once
            # (at its first cell) regardless of fill/text rules below.
            if merge_rng.min_col == col:
                text = ws.cell(row=merge_rng.min_row, column=merge_rng.min_col).value
                end_col = min(merge_rng.max_col, header.column_for_day(days))
                ref = f"{sheet_name}!{cell_ref(berth_row, col)}:{cell_ref(berth_row, end_col)}"
                if isinstance(text, str) and text.strip():
                    spans.append(Span(berth_label, col, end_col, text.strip(), ref))
                else:
                    issues.append(
                        SpanIssue(
                            "ORPHAN_FILL",
                            f'Merged range with no text on "{berth_label}"',
                            ref,
                            berth_label=berth_label,
                            start_col=col,
                            end_col=end_col,
                        )
                    )
            next_day = header.day_for_column(merge_rng.max_col + 1)
            day = next_day if next_day is not None else days + 1
            continue

        if not _is_booking_cell(ws, berth_row, col, config):
            day += 1
            continue

        run_cells: list[tuple[int, str | None]] = []
        d = day
        while d <= days:
            c = header.column_for_day(d)
            if merge_lookup.get((berth_row, c)) is not None:
                break
            if not _is_booking_cell(ws, berth_row, c, config):
                break
            value = ws.cell(row=berth_row, column=c).value
            text = value.strip() if isinstance(value, str) and value.strip() else None
            run_cells.append((c, text))
            d += 1
        day = d

        run_spans, run_issues = _spans_from_run(run_cells, berth_label, sheet_name, berth_row)
        spans.extend(run_spans)
        issues.extend(run_issues)

    return spans, issues


@dataclass
class UnlabeledEntry:
    berth_above: str
    text: str
    source_ref: str


# How many rows past the last recognized berth row to scan for stray text
# before giving up (SPEC.md section 7.3 rule 7). Bounds a runaway scan if a
# block's true end (next weekday row) isn't found for some reason; every
# observed gap in the real workbook is well under this.
UNLABELED_SCAN_LIMIT = 15


def find_unlabeled_entries(
    ws: Worksheet, header: MonthHeader, berth_rows: list[tuple[int, str]], sheet_name: str
) -> list[UnlabeledEntry]:
    """Text typed into a row that has no berth label of its own — e.g. a
    vessel name in a blank row under the last defined berth. Scans the gap
    between the last berth row and the next weekday row / month header."""
    from importer.layout import MONTH_LABEL_RE

    if not berth_rows:
        return []
    last_berth_row, last_berth_label = berth_rows[-1]

    entries: list[UnlabeledEntry] = []
    for row in range(last_berth_row + 1, last_berth_row + 1 + UNLABELED_SCAN_LIMIT):
        if row > ws.max_row:
            break
        label = ws.cell(row=row, column=1).value
        if is_weekday_row(ws, row):
            break
        if isinstance(label, str) and (label.strip() == "" or MONTH_LABEL_RE.match(label.strip())):
            break
        if isinstance(label, str) and label.strip():
            break
        for day in range(1, header.days_in_month() + 1):
            col = header.column_for_day(day)
            value = ws.cell(row=row, column=col).value
            if isinstance(value, str) and value.strip():
                entries.append(
                    UnlabeledEntry(
                        berth_above=last_berth_label,
                        text=value.strip(),
                        source_ref=f"{sheet_name}!{cell_ref(row, col)}",
                    )
                )
    return entries


def berth_rows_for_block(ws: Worksheet, header: MonthHeader) -> list[tuple[int, str]]:
    """Rows from `header.first_berth_row` down to the next weekday row,
    month-header row, or blank row — i.e. the berth rows belonging to this
    one month block. Returns (row, raw column-A label)."""
    from importer.layout import MONTH_LABEL_RE

    rows: list[tuple[int, str]] = []
    row = header.first_berth_row
    while row <= ws.max_row:
        label = ws.cell(row=row, column=1).value
        if label is None:
            break
        if not isinstance(label, str):
            break
        stripped = label.strip()
        if is_weekday_row(ws, row) or MONTH_LABEL_RE.match(stripped):
            break
        rows.append((row, stripped))
        row += 1
    return rows
