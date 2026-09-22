"""Parses the Science and Yachts tabs for vessel details (SPEC.md section
7.5). These tabs are loosely structured: a vessel-prefix cell in column A
starts a record; every cell until the next such cell (any column) holds
operator/contact/phone/email/notes in no fixed layout. Matched to schedule
vessels via normalized_key (section 7.4/7.6); best-effort only, exactly as
the spec asks — contact details are dumped into the vessel's notes with a
"verify" disclaimer rather than precision-matched to a specific person.
"""

import re
from dataclasses import dataclass, field

from openpyxl.worksheet.worksheet import Worksheet

from importer.normalize import VesselMention, parse_vessel_mention
from importer.ruleset import ImporterConfig

LENGTH_RE = re.compile(r"(\d+(?:\.\d+)?)\s*'")
LOA_RE = re.compile(r"LOA:\s*(\d+(?:\.\d+)?)\s*'", re.IGNORECASE)
DRAFT_RE = re.compile(r"Draft:\s*(\d+(?:\.\d+)?)\s*'", re.IGNORECASE)
EMAIL_RE = re.compile(r"[\w.+-]+@[\w-]+\.[\w.-]+")
PHONE_RE = re.compile(r"Cell:\s*[\d-]+", re.IGNORECASE)
CAPTAIN_RE = re.compile(r"Capt\.\s*[A-Za-z][A-Za-z .'-]*")

CONTACT_DISCLAIMER = "Contact details imported automatically — verify."


@dataclass
class VesselDetailRecord:
    mention: VesselMention
    loa_ft: float | None
    draft_ft: float | None
    organization_name: str | None
    captain_names: list[str] = field(default_factory=list)
    phones: list[str] = field(default_factory=list)
    emails: list[str] = field(default_factory=list)
    conflicting_length: tuple[float, float] | None = None  # (used, discarded)
    source_ref: str = ""


def _record_rows(ws: Worksheet, start_row: int, next_start_row: int | None) -> range:
    end_row = (next_start_row - 1) if next_start_row is not None else ws.max_row
    return range(start_row, end_row + 1)


def _all_cell_texts(ws: Worksheet, rows: range) -> list[tuple[int, int, str]]:
    out = []
    for row in rows:
        for col in range(1, ws.max_column + 1):
            v = ws.cell(row=row, column=col).value
            if isinstance(v, str) and v.strip():
                out.append((row, col, v.strip()))
    return out


def parse_vessel_detail_tab(
    ws: Worksheet, sheet_name: str, config: ImporterConfig
) -> list[VesselDetailRecord]:
    # Find every record-start row (column A holds a vessel-prefix mention).
    starts: list[tuple[int, VesselMention]] = []
    for row in range(1, ws.max_row + 1):
        value = ws.cell(row=row, column=1).value
        if not isinstance(value, str):
            continue
        mention = parse_vessel_mention(value, config)
        if mention is not None:
            starts.append((row, mention))

    records: list[VesselDetailRecord] = []
    for i, (start_row, mention) in enumerate(starts):
        next_start_row = starts[i + 1][0] if i + 1 < len(starts) else None
        rows = _record_rows(ws, start_row, next_start_row)
        cells = _all_cell_texts(ws, rows)

        # A trailing length on the vessel-prefix cell itself, e.g.
        # "M/Y Grey Horizon 46'" -> mention.name already has "46'" in it
        # (parse_vessel_mention only strips the prefix). Pull it back out.
        name_length_match = LENGTH_RE.search(mention.name)
        prefix_cell_length = float(name_length_match.group(1)) if name_length_match else None
        clean_name = LENGTH_RE.sub("", mention.name).strip()
        clean_mention = VesselMention(
            raw_text=mention.raw_text, prefix=mention.prefix, name=clean_name
        )

        loa_candidates: list[float] = []
        draft_candidates: list[float] = []
        organization_name: str | None = None
        captains: list[str] = []
        phones: list[str] = []
        emails: list[str] = []

        for row, col, text in cells:
            if row == start_row and col == 1:
                continue  # the vessel-prefix cell itself, already handled

            loa_match = LOA_RE.search(text)
            draft_match = DRAFT_RE.search(text)
            if loa_match:
                loa_candidates.append(float(loa_match.group(1)))
            if draft_match:
                draft_candidates.append(float(draft_match.group(1)))
            if loa_match or draft_match:
                continue

            email_match = EMAIL_RE.search(text)
            if email_match:
                emails.append(email_match.group(0))
                continue
            if PHONE_RE.search(text):
                phones.append(text)
                continue
            captain_match = CAPTAIN_RE.search(text)
            if captain_match:
                captains.append(captain_match.group(0).strip())
                continue

            # First remaining plain-text cell in the record (not the vessel
            # cell, a length, a phone, an email, or a captain) is taken as
            # the operating organization — matches every example in the
            # real data, where it's the cell right after the vessel name.
            if organization_name is None:
                organization_name = text

        loa_ft = loa_candidates[0] if loa_candidates else prefix_cell_length
        conflicting_length = None
        if (
            loa_candidates
            and prefix_cell_length is not None
            and loa_candidates[0] != prefix_cell_length
        ):
            conflicting_length = (loa_candidates[0], prefix_cell_length)

        records.append(
            VesselDetailRecord(
                mention=clean_mention,
                loa_ft=loa_ft,
                draft_ft=draft_candidates[0] if draft_candidates else None,
                organization_name=organization_name,
                captain_names=captains,
                phones=phones,
                emails=emails,
                conflicting_length=conflicting_length,
                source_ref=f"{sheet_name}!A{start_row}",
            )
        )
    return records
